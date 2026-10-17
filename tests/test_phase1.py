"""Automated verification test for Phase 1 components."""

import io
import pytest
from pathlib import Path
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from agentic_ai.parsers import (
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_document,
)
from agentic_ai.schemas import (
    ResumeEvaluation,
    BulletOptimization,
    SkillGapReport,
    ScoringRubric,
)
from agentic_ai.evaluator import mock_evaluate_resume, evaluate_resume
from agentic_ai.sample_data import SAMPLE_RESUME, SAMPLE_JOB_DESCRIPTION


def test_docx_parser():
    """Test extracting text from an in-memory DOCX document."""
    doc = docx.Document()
    doc.add_paragraph("First paragraph of resume.")
    doc.add_paragraph("Second paragraph with Python and React skills.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill: Machine Learning"
    table.cell(0, 1).text = "Level: Advanced"
    table.cell(1, 0).text = "Skill: Docker"
    table.cell(1, 1).text = "Level: Intermediate"

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    extracted = extract_text_from_docx(bio)
    assert "First paragraph of resume." in extracted
    assert "Python and React" in extracted
    assert "Machine Learning" in extracted
    assert "Docker" in extracted


def test_pdf_parser():
    """Test extracting text from an in-memory generated PDF using pdfplumber."""
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Jane Doe - Senior AI Engineer", styles["Title"]),
        Paragraph("Experienced in PyTorch, LangChain, and FastAPI microservices.", styles["Normal"]),
    ]
    doc.build(story)
    bio.seek(0)

    extracted = extract_text_from_pdf(bio)
    assert "Jane Doe" in extracted
    assert "PyTorch" in extracted
    assert "FastAPI" in extracted


def test_parse_document_dispatcher():
    """Test the universal dispatcher with plain text and bytes."""
    sample_text = "This is a plain text resume content."
    res = parse_document(sample_text.encode("utf-8"), filename="resume.txt")
    assert res == sample_text


def test_schemas_validation():
    """Test that the Pydantic schemas enforce type safety and bounds."""
    evaluation = mock_evaluate_resume(SAMPLE_RESUME, SAMPLE_JOB_DESCRIPTION)
    assert isinstance(evaluation, ResumeEvaluation)
    assert 0 <= evaluation.overall_match_score <= 100
    assert len(evaluation.skills_alignment.matched_skills) > 0
    assert len(evaluation.bullet_rewrites) >= 3

    # Ensure bullet optimizations have necessary fields
    for bullet in evaluation.bullet_rewrites:
        assert bullet.original != ""
        assert bullet.improved != ""
        assert bullet.explanation != ""


def test_evaluator_execution():
    """Test that evaluate_resume returns a valid evaluation object."""
    report, provider = evaluate_resume(
        SAMPLE_RESUME, SAMPLE_JOB_DESCRIPTION, force_provider="mock"
    )
    assert report.overall_match_score > 0
    assert provider == "mock"


if __name__ == "__main__":
    print("Running Phase 1 verification tests...")
    test_docx_parser()
    print(" [x] DOCX parser test passed.")
    test_pdf_parser()
    print(" [x] PDF parser (pdfplumber) test passed.")
    test_parse_document_dispatcher()
    print(" [x] Universal dispatcher test passed.")
    test_schemas_validation()
    print(" [x] Pydantic schemas validation passed.")
    test_evaluator_execution()
    print(" [x] Baseline evaluator execution test passed.")
    print("All Phase 1 tests passed successfully!")
