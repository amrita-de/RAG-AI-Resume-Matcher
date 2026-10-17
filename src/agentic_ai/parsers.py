"""Document parsing module for resumes and job descriptions.

Supports PDF (via pdfplumber for resilient multi-column text extraction),
DOCX (via python-docx for paragraph and table extraction), and plain text.
"""

from pathlib import Path
from typing import Union, BinaryIO
import io
import re
import pdfplumber
import docx


def clean_extracted_text(text: str) -> str:
    """Normalize whitespace and strip excessive blank lines from extracted text."""
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    cleaned = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


def extract_text_from_pdf(file_source: Union[str, Path, bytes, BinaryIO]) -> str:
    """Extract clean text from a PDF file using pdfplumber.

    Handles layout and multi-column formats better than legacy PyPDF2.
    """
    text_chunks = []

    if isinstance(file_source, bytes):
        stream = io.BytesIO(file_source)
    elif isinstance(file_source, (str, Path)):
        stream = open(file_source, "rb")
    else:
        stream = file_source

    try:
        with pdfplumber.open(stream) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text(layout=True)
                if not page_text:
                    # Fallback to simple extraction if layout mode yielded empty text
                    page_text = page.extract_text()
                if page_text:
                    text_chunks.append(clean_extracted_text(page_text))
    finally:
        if isinstance(file_source, (str, Path)) and not stream.closed:
            stream.close()

    return "\n\n".join(text_chunks).strip()


def extract_text_from_docx(file_source: Union[str, Path, bytes, BinaryIO]) -> str:
    """Extract clean text from a .docx file including paragraphs and tables."""
    if isinstance(file_source, bytes):
        stream = io.BytesIO(file_source)
    else:
        stream = file_source

    doc = docx.Document(stream)
    text_chunks = []

    # Extract all paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_chunks.append(para.text.strip())

    # Extract table cells (resumes often use tables for skills or two-column sections)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                # Deduplicate identical adjacent cells (due to merged cells)
                unique_texts = []
                for t in row_texts:
                    if not unique_texts or unique_texts[-1] != t:
                        unique_texts.append(t)
                text_chunks.append(" | ".join(unique_texts))

    return "\n".join(text_chunks).strip()


def parse_document(file_source: Union[str, Path, bytes, BinaryIO], filename: str = "") -> str:
    """Universal parser dispatcher that extracts raw text based on file format.

    Args:
        file_source: Path to file, byte contents, or file-like stream.
        filename: Optional filename used to infer format if file_source is bytes/stream.

    Returns:
        Extracted plain text content.
    """
    ext = ""
    if isinstance(file_source, (str, Path)):
        ext = Path(file_source).suffix.lower()
    elif filename:
        ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_source)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_source)
    elif ext in [".txt", ".md"]:
        if isinstance(file_source, bytes):
            return file_source.decode("utf-8", errors="ignore").strip()
        elif isinstance(file_source, (str, Path)):
            return Path(file_source).read_text(encoding="utf-8", errors="ignore").strip()
        else:
            return file_source.read().decode("utf-8", errors="ignore").strip()
    else:
        # Attempt fallback to UTF-8 text decode
        if isinstance(file_source, bytes):
            return file_source.decode("utf-8", errors="ignore").strip()
        elif isinstance(file_source, (str, Path)):
            return Path(file_source).read_text(encoding="utf-8", errors="ignore").strip()
        raise ValueError(f"Unsupported file format '{ext}' for document parsing.")
